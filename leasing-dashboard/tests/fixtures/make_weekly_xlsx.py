"""
テスト用の週報 Excel（架空データ）を作る。実在の物件・人物・会社とは関係ない。

  python3 tests/fixtures/make_weekly_xlsx.py

様式 A〜D の見出し・節の構成だけを実際の様式に合わせ、値はすべて架空。
数式はキャッシュ値なしで保存される（実際の週報と同じく、計算結果がファイルにない状態を再現する）。
"""
import datetime as dt
import os

from openpyxl import Workbook

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'weekly')
D = dt.datetime


def serial(d):
    """Excel のシリアル値（日付書式なしの数値で入っている列の再現）"""
    return (d - dt.datetime(1899, 12, 30)).days


def format_a(path, blank_status=False):
    wb = Workbook()
    ws = wb.active
    ws.title = '週報'
    ws['B2'] = '週報'
    ws['B3'] = '基準日'
    ws['C3'] = D(2026, 1, 11)
    ws['B4'] = '集計'; ws['C4'] = '坪数'; ws['D4'] = '部屋数'; ws['M4'] = '稼働率（全体）'
    for i, lab in enumerate(['解約予定', '空室', '原復中', '申込', '契約済', '募集止め', 'FR']):
        r = 5 + i
        ws.cell(r, 2, lab)
        ws.cell(r, 4, '=COUNTIF($B$16:$B$40,B%d)' % r)
    ws['N5'] = '=1-SUM(C6:C8)/基本情報!$C$4'
    heads = {'B': '現況', 'C': '部屋番号', 'D': '間取り', 'E': '貸室面積\n (坪)', 'F': '貸室面積\n(㎡)', 'G': '解約日', 'H': '退去日',
             'I': '空室日数', 'J': 'Downtime\n(day)', 'K': '原復完了', 'L': '工事期間', 'M': '従前賃料(共込）', 'N': '予算賃料\n（共込）',
             'O': '募集賃料', 'P': '管理費', 'Q': '賃料計', 'R': '坪単価', 'S': '敷金', 'T': '礼金', 'U': 'FR', 'V': 'AD', 'W': '申込日',
             'X': '法人/個人', 'Y': '保証会社', 'AA': '賃料発生日', 'AB': '賃料', 'AC': '管理費', 'AJ': '契約形態'}
    ws['B14'] = '空室状況'; ws['M14'] = '募集条件'; ws['W14'] = '契約条件'
    for col, h in heads.items():
        ws[col + '15'] = h
    rows = [
        # 現況, 部屋, 解約日, 退去日, 原復完了, 募集賃料, 管理費, 敷金, 礼金, FR, AD, 申込日, 賃料発生日, 賃料
        ('契約済', 101, None, None, None, None, None, None, None, None, None, None, '2024/04/01', 100000),
        ('空室', 102, D(2025, 10, 1), D(2025, 10, 3), D(2025, 10, 20), 120000, 5000, 0, 1, 0, 1, None, None, None),
        ('原復中', 103, D(2025, 12, 20), D(2025, 12, 25), D(2026, 1, 20), 118000, 5000, 0, 1, 0, 1, None, None, None),
        ('申込', 201, D(2025, 11, 1), D(2025, 11, 1), D(2025, 11, 15), 121000, 5000, 0, 0, 1, 2, D(2026, 1, 8), None, None),
        ('解約予定', 202, D(2026, 2, 28), None, None, None, None, None, None, None, None, None, None, None),
        ('契約済', 203, D(2025, 9, 1), D(2025, 9, 1), D(2025, 9, 20), 119000, 5000, 0, 0, 0, 1, D(2025, 12, 10), D(2026, 2, 1), 119000),
        ('募集止め', 301, D(2025, 9, 1), D(2025, 9, 2), None, None, None, None, None, None, None, None, None, None),
        ('FR', 302, None, None, None, None, None, None, None, None, None, None, '2026/01/01', 99000),
        ('契約済', 303, None, None, None, None, None, None, None, None, None, None, '2023/06/15', 101000),
    ]
    if blank_status:
        rows[0] = (None,) + rows[0][1:]
    for i, r in enumerate(rows):
        n = 16 + i
        st, room, cancel, moveout, restored, asking, fee, dep, key, fr, ad, appd, rent_start, rent = r
        if st: ws['B%d' % n] = st
        ws['C%d' % n] = room
        ws['F%d' % n] = 30.5
        for col, v in (('G', cancel), ('H', moveout), ('K', restored), ('O', asking), ('P', fee), ('S', dep), ('T', key),
                       ('U', fr), ('V', ad), ('W', appd), ('AA', rent_start), ('AB', rent)):
            if v is not None:
                ws['%s%d' % (col, n)] = v
        ws['I%d' % n] = '=IF(OR(B%d="空室",B%d="原復中")=TRUE,IF(G%d="","",基本情報!$C$1-G%d),"")' % (n, n, n, n)
        ws['Q%d' % n] = '=IF(O%d="","",O%d+P%d)' % (n, n, n)
        if st in ('契約済', 'FR'):
            ws['X%d' % n] = '個人'
            ws['Y%d' % n] = '架空保証株式会社'
    ws['B27'] = 'リーシング状況'
    ws['Z27'] = 'リーシング活動・条件変更など'
    ws['E28'] = '累計'
    for i, d in enumerate([D(2025, 12, 28), D(2026, 1, 4), D(2026, 1, 11)]):
        ws.cell(28, 6 + i, d)
    ws['AA28'] = '実施日'; ws['AB28'] = '変更内容等'
    for i, (lab, vals) in enumerate([('問合せ', (2, 3, 4)), ('案内', (1, 2, 3)), ('申込', (0, 0, 1)), ('キャンセル', (0, 0, 0))]):
        ws.cell(29 + i, 2, lab)
        for j, v in enumerate(vals):
            ws.cell(29 + i, 6 + j, v)
    ws['AA29'] = D(2025, 12, 26); ws['AB29'] = '架空: 仲介会社訪問'
    ws['AA30'] = D(2026, 1, 9); ws['AB30'] = '架空: 102号室の募集条件を見直し'
    ws['B35'] = '申込進捗状況'
    ws['B36'] = '申込日'; ws['C36'] = '申込住戸'; ws['D36'] = '進捗'; ws['F36'] = '内容'
    ws['B37'] = D(2026, 1, 8); ws['C37'] = 201; ws['D37'] = '審査中'; ws['F37'] = '個人・架空勤務先株式会社勤務'
    info = wb.create_sheet('基本情報')
    info['A2'] = '物件名'; info['C2'] = 'テスト物件A'
    info['C1'] = '=週報!C3'
    info['A3'] = 'Unit 数'; info['C3'] = '=COUNT(A8:A1048575)'
    info['A4'] = 'NRA (坪)'; info['C4'] = '=SUM(D8:D1048576)'
    for col, h in zip('ABCDE', ['部屋番号', '用途', '間取り', '貸室面積 (坪)', '貸室面積 (㎡)']):
        info[col + '7'] = h
    for i, room in enumerate([101, 102, 103, 201, 202, 203, 301, 302, 303]):
        info.cell(8 + i, 1, room); info.cell(8 + i, 2, '住戸'); info.cell(8 + i, 3, ' 1K'); info.cell(8 + i, 5, 30.5)
    wb.create_sheet('業者営業')['B1'] = '訪問日'
    wb.save(path)


def format_b(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '週報'
    ws['A1'] = '営業活動報告書'; ws['R1'] = '対象期間：'; ws['T1'] = '=A10'; ws['W1'] = '=T1+6'
    ws['A4'] = 'テスト物件B'; ws['L4'] = '総戸数'; ws['N4'] = 20; ws['P4'] = '申込～契約前'; ws['R4'] = '=COUNTIF(N32:O35,"審査中")'
    ws['L5'] = '契約中'; ws['N5'] = '=N4-R6'; ws['P5'] = '解約予定'
    ws['L6'] = '戸数稼働率'; ws['N6'] = '=N5/N4'; ws['P6'] = '空室'
    ws['A8'] = '【案内状況】'; ws['L8'] = '【募集条件】'
    for col, h in zip('ACEGI', ['日付', '曜日', '案内件数', '申込件数', '契約件数']):
        ws[col + '9'] = h
    ws['A10'] = D(2026, 1, 5)
    view = [1, 0, 2, 0, 0, 1, 0]
    app = [0, 0, 1, 0, 0, 0, 0]
    for i, w in enumerate('月火水木金土日'):
        ws.cell(10 + i, 3, w); ws.cell(10 + i, 5, view[i]); ws.cell(10 + i, 7, app[i]); ws.cell(10 + i, 9, 0)
    ws['A17'] = '週間合計'; ws['E17'] = '=SUM(E10:F16)'
    ws['A20'] = '【募集中・解約予定住戸】'; ws['L20'] = '【内覧履歴】'
    for col, h in zip(['A', 'B', 'C', 'D', 'E', 'G', 'I', 'L', 'M', 'O', 'Q', 'S'], ['号室', '状況', 'リノベ', '㎡数', '解約日', '前回賃料', '募集賃料', '内覧日', '住戸', '状況', '仲介業者', 'コメント']):
        ws[col + '21'] = h
    for i, r in enumerate([(501, '空室', '済', 40.5, D(2025, 11, 30), 150000, 165000), (302, '申込', None, 30.2, D(2025, 12, 15), 120000, 125000),
                           (205, '解約予定', None, 25.0, D(2026, 2, 15), 98000, None)]):
        for col, v in zip(['A', 'B', 'C', 'D', 'E', 'G', 'I'], r):
            if v is not None:
                ws['%s%d' % (col, 22 + i)] = v
    ws['L22'] = D(2026, 1, 6); ws['M22'] = 302; ws['O22'] = '申込'; ws['Q22'] = '架空不動産'; ws['S22'] = '架空法人申込・年収800万'
    ws['A29'] = '【条件変更等】'
    ws['A30'] = '実施日'; ws['B30'] = '変更内容等'; ws['L30'] = '【進捗状況】'
    ws['A31'] = D(2025, 12, 1); ws['B31'] = '架空: 定期借家の条件を追加'
    ws['A32'] = D(2026, 1, 7); ws['B32'] = '501号室：礼金0ヶ月'
    for col, h in zip(['L', 'M', 'N', 'P', 'R', 'T'], ['申込日', '住戸', '進捗状況', '契約賃料（管込）', '開始日', 'コメント']):
        ws[col + '31'] = h
    ws['L32'] = D(2026, 1, 6); ws['M32'] = 302; ws['N32'] = '審査中'; ws['P32'] = 125000; ws['T32'] = '架空: 保証会社審査中'
    wb.create_sheet('スタッキング')['B1'] = '【賃料表】'
    wb.save(path)


def c_week(ws, title, start, end, report, summary, apps, notices, cancels):
    ws['B1'] = title
    ws['B3'] = start.year; ws['C3'] = start.month; ws['D3'] = '月'; ws['E3'] = start.day; ws['F3'] = '日'; ws['H3'] = '～'
    for col, h in zip(['B', 'C', 'D', 'F', 'H', 'J', 'K', 'M', 'O'], ['報告日', '空室数', '募集戸数', '反響数', '内見数', '新規申込', '申込中\n契約予定', '申込キャンセル数', '契約数']):
        ws[col + '5'] = h
    ws['B6'] = serial(report)
    for col, v in zip(['C', 'D', 'F', 'H', 'J', 'K', 'M'], summary):
        ws[col + '6'] = v
    ws['O6'] = '=COUNTA(B26:B40)'
    ws['B17'] = '＜申込＞'
    for col, h in zip(['B', 'C', 'F', 'H', 'J', 'R', 'S'], ['号室', '申込者名', '申込日', '契約予定日', '賃料', '属性', '年収']):
        ws[col + '18'] = h
    for i, (room, d) in enumerate(apps):
        ws['B%d' % (19 + i)] = room; ws['C%d' % (19 + i)] = '個人'; ws['F%d' % (19 + i)] = serial(d); ws['R%d' % (19 + i)] = '会社員'; ws['S%d' % (19 + i)] = 600
    ws['B24'] = '＜契約＞'
    ws['B30'] = '＜解約予告＞'
    for col, h in zip(['B', 'C', 'F', 'H'], ['号室', '契約者名', '解約日', '賃料']):
        ws[col + '31'] = h
    for i, (room, d) in enumerate(notices):
        ws['B%d' % (32 + i)] = room; ws['C%d' % (32 + i)] = '架空商事株式会社'; ws['F%d' % (32 + i)] = serial(d); ws['H%d' % (32 + i)] = 95000
    ws['B36'] = '【その他】ｷｬﾝｾﾙ'
    for i, (room, d) in enumerate(cancels):
        ws['B%d' % (37 + i)] = room; ws['C%d' % (37 + i)] = '個人'; ws['F%d' % (37 + i)] = serial(d); ws['H%d' % (37 + i)] = 92000; ws['N%d' % (37 + i)] = '架空の理由'


def format_c(path):
    wb = Workbook()
    wb.active.title = 'スタッキング'
    cur = wb.create_sheet('現募集中')
    cur['R2'] = '報告日'; cur['S2'] = serial(D(2026, 1, 12))
    for col, h in zip(['B', 'C', 'E', 'F', 'G', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'S'],
                      ['号室', '面積', '間取り', '賃料', '共益費', 'AD', '敷金', '礼金', 'フリーレント', '募集賃料', '解約日', '原状回復完了日', '募集開始日', '備考']):
        cur[col + '3'] = h
    cur['C4'] = '㎡'; cur['E4'] = 'タイプ'; cur['M4'] = '提案'
    rows = [(None, 101, 90000, None, None, None), ('空室', 102, 91000, D(2025, 11, 20), D(2025, 12, 5), 1), ('申込', 103, 92000, D(2025, 12, 1), D(2026, 1, 20), 1),
            (None, 104, 93000, D(2026, 1, 31), None, None), ('退去予定', 201, 95000, None, None, None), ('契約中', 202, 96000, None, None, None)]
    for i, (st, room, rent, cancel, restored, ad) in enumerate(rows):
        n = 5 + i
        if st: cur['A%d' % n] = st
        cur['B%d' % n] = room; cur['F%d' % n] = rent; cur['G%d' % n] = 5000
        if cancel: cur['N%d' % n] = serial(cancel)
        if restored: cur['O%d' % n] = serial(restored)
        if ad is not None: cur['I%d' % n] = ad; cur['J%d' % n] = 0; cur['K%d' % n] = 1; cur['L%d' % n] = 0
    ov = wb.create_sheet('部屋概要')
    for col, h in zip(['B', 'C', 'E', 'F'], ['号室', '面積', '間取り', '賃料']):
        ov[col + '3'] = h
    ov['C4'] = '㎡'; ov['D4'] = '坪'; ov['E4'] = 'タイプ'
    for i, room in enumerate([101, 102, 103, 104, 105, 201, 202, 203, 204, 205]):
        ov['B%d' % (5 + i)] = room; ov['C%d' % (5 + i)] = 25.0; ov['E%d' % (5 + i)] = '1K'
    c_week(wb.create_sheet('2026.1.5-2026.1.11'), 'テスト物件C', D(2026, 1, 5), D(2026, 1, 11), D(2026, 1, 12), (3, 3, 4, 2, 1, 1, 1),
           apps=[(103, D(2026, 1, 8))], notices=[(105, D(2025, 11, 20)), (201, D(2026, 2, 28))], cancels=[(102, D(2025, 12, 20)), (104, D(2026, 1, 9))])
    c_week(wb.create_sheet('2025.12.29-2026.1.4'), 'テスト物件C', D(2025, 12, 29), D(2026, 1, 4), D(2026, 1, 5), (3, 3, 2, 1, 0, 0, 0),
           apps=[], notices=[(105, D(2025, 11, 20))], cancels=[(102, D(2025, 12, 20))])
    wb.save(path)


def format_d(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'List'
    ws['A1'] = 'Action Points List'; ws['A2'] = 'Asset Manager:'; ws['C2'] = 'Date:'
    for col, h in zip('ABCDEF', ['No.', 'Property name', 'Target Due Date', 'Completion', 'Priority', 'Status']):
        ws[col + '4'] = h
    for col, h in zip('CFGHIJK', ['Quote/Proposal', 'Status', 'Category', 'Action', 'Memo', 'Latest Updated', '架空担当メモ']):
        ws[col + '5'] = h
    items = [(1, 'テスト物件A', D(2026, 1, 5), '高', '未完了', 'EV点検', '見積待ち', '居住者 架空太郎 から連絡', '架空: 見積依頼済', '架空メモ'),
             (None, 'テスト物件C', D(2026, 2, 1), '中', '完了（報告済）', '共用灯交換', '交換済み', None, '架空: 1/6 交換完了', None),
             (None, '未登録の物件', D(2026, 1, 20), '低', '未完了', '植栽', '剪定', None, None, None)]
    for i, it in enumerate(items):
        for col, v in zip('ABCEFGHIJK', it):
            if v is not None:
                ws['%s%d' % (col, 6 + i)] = v
    ws['A20'] = '=COUNTA(B6:B19)'
    dd = wb.create_sheet('Dropdown')
    dd['A2'] = 'Priority'; dd['C3'] = 'PJ No.'; dd['D3'] = 'PJ Name'; dd['E3'] = 'Property Name'
    dd['C4'] = 103; dd['E4'] = 'テスト物件C'
    wb.save(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    format_a(os.path.join(OUT, '【テスト物件A】週報.xlsx'))
    format_a(os.path.join(OUT, '..', 'weekly_variants', 'A_blank_status.xlsx') if os.makedirs(os.path.join(OUT, '..', 'weekly_variants'), exist_ok=True) is None else None, blank_status=True)
    format_b(os.path.join(OUT, 'テスト物件B_営業活動報告.xlsx'))
    format_c(os.path.join(OUT, 'リーシング報告書_(テスト物件C).xlsx'))
    format_d(os.path.join(OUT, 'ActionPoints_test.xlsx'))
    unknown = Workbook()
    unknown.active['A1'] = '架空の別様式'
    unknown.save(os.path.join(OUT, 'unknown_format.xlsx'))
    with open(os.path.join(OUT, 'broken.xlsx'), 'w', encoding='utf-8') as f:
        f.write('これは Excel ではありません')
    with open(os.path.join(OUT, '..', 'weekly_master.csv'), 'w', encoding='utf-8') as f:
        f.write('project_number,project_name,property_name_ja,property_name_en,spc,am,pm,master_source\n'
                '101,Test A,テスト物件A,Test A,架空SPC,架空AM,架空PM,test\n'
                '102,Test B,テスト物件B,Test B,架空SPC,架空AM,架空PM,test\n'
                '103,Test C,テスト物件C,Test C,架空SPC,架空AM,架空PM,test\n')
    print('written to', OUT)


if __name__ == '__main__':
    main()
